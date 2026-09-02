#!/usr/bin/env python3
"""
Run scout_selection-style medoid analysis across cutoff dates and BETA values,
then export an Excel report with medoid lineages, fitness scores, and distances.

Derived from scripts/scout_selection.ipynb.
"""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from Bio import SeqIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent
sys.path.insert(0, str(SCRIPT_DIR))

from Helpers import (  # noqa: E402
    add_fitness_scores_pyro,
    add_mutations_column,
    compute_fitness_score_pyro,
    compute_mutations,
    getReference,
    read_pickle,
    slice_by_disjoint_windows,
    split_circulating_and_future,
)

ROOT_PATH_DEFAULT = "/data/pgangwar/vaccinescout/data/"
CUTOFF_DATES_DEFAULT = ["2021-06-01", "2022-06-01", "2023-06-01", "2024-06-01"]
BETAS_DEFAULT = [0, 10, 20, 30, 40, 50, 60, 70, 80, 90, 100]
FITNESS_COLUMN = "fitness_mean"
MIN_TIP_COUNT = 5
CIRCULATION_DAYS = 30
DISJOINT_WINDOWS = ["6M_1", "6M_2"]
K_VALUES = [1, 2, 3, 4, 5, 6, 7, 8]
CURRENT_DISTANCE_BETA = 60
PYRO_VER = "Jan_4_24"
BLOCK_COLS = 9
BLOCK_GUTTER = 1

FDA_STRAINS_BY_CUTOFF = {
    "2021-06-01": "B",
    "2022-06-01": "BA.4",
    "2023-06-01": "XBB.1.5",
    "2024-06-01": "KP.2",
}

PFIZER_INN_BY_CUTOFF = {
    "2021-06-01": "Tozinameran",
    "2022-06-01": "Famtozinameran",
    "2023-06-01": "Raxtozinameran",
    "2024-06-01": "Cemivameran",
}

PFIZER_FASTA_BY_CUTOFF = {
    "2021-06-01": "tozinameran_2021_aa_modified.fa",
    "2022-06-01": "famtozinameran_2022_aa_nodel_modified.fa",
    "2023-06-01": "raxtozinameran_2023_aa_nodel_modified.fa",
    "2024-06-01": "cemivameran_2024_aa_noindel_modified.fa",
}


EXCEL_COLUMNS = [
    "Weights Function",
    "Neutralization Scores Function",
    "Medoids",
    "Lineages",
    "CoVFit Fitness",
    "PyR0 Fitness",
    "Distance to Current",
    "Distance to Future (0-6 months)",
    "Distance to Future (6-12 months)",
]

# Reference-vaccine rows use the first three medoid-block columns for the name only.
FDA_EXCEL_COLUMNS = [
    "FDA Variant",
    "",
    "",
    "Lineages",
    "CoVFit Fitness",
    "PyR0 Fitness",
    "Distance to Current",
    "Distance to Future (0-6 months)",
    "Distance to Future (6-12 months)",
]

PFIZER_EXCEL_COLUMNS = [
    "Pfizer Vaccine",
    "",
    "",
    "Lineages",
    "CoVFit Fitness",
    "PyR0 Fitness",
    "Distance to Current",
    "Distance to Future (0-6 months)",
    "Distance to Future (6-12 months)",
]

ORACLE_INTERVAL_LABELS = {
    "6M_1": "0-6 Months",
    "6M_2": "6-12 Months",
}

ORACLE_EXCEL_COLUMNS = [
    "Oracle Fit Slice",
    "",
    "Medoids",
    "Lineages",
    "CoVFit Fitness",
    "PyR0 Fitness",
    "Distance to Current",
    "Distance to Future (0-6 months)",
    "Distance to Future (6-12 months)",
]


def load_data(file_path: Path) -> dict:
    df = pd.read_parquet(file_path)
    return dict(zip(df["key"], df["value"]))


def exp_fitness_modifier(fitness_scores, beta):
    """Fitness modifier: e ** (beta * fitness)."""
    s_arr = np.asarray(fitness_scores, dtype=np.float32)
    return np.power(np.e, beta * s_arr)


def effective_weights(weights, fitness_scores=None, beta=None):
    """
    Effective tip weights w * e ** (beta * fitness), computed in log-space and
    rescaled so the largest weight is 1.0 (overflow-safe for any beta).

    Working in log-space avoids the e ** (beta * fitness) overflow:

        log_w_eff = ln(w) + beta * fitness
        w_eff     = exp(log_w_eff - max(log_w_eff))      # in (0, 1]

    Subtracting the max multiplies every weight by the same positive constant
    e ** (-max), so it does not change any argmin / relative cost comparison in
    the PAM selection; it only fixes the numerical scale.

    Returns float64 so the weighted cost sums and the swap-acceptance epsilon
    stay meaningful (the weights vector is 1-D, so the memory cost is trivial).
    """
    w = np.asarray(weights, dtype=np.float64)
    if fitness_scores is None or beta is None:
        return w

    f = np.asarray(fitness_scores, dtype=np.float64)
    # ln(0) -> -inf yields a clean zero effective weight after exp; guard against
    # the degenerate all-zero case where max would be -inf.
    with np.errstate(divide="ignore"):
        log_w = np.where(w > 0, np.log(w), -np.inf)
    log_w_eff = log_w + float(beta) * f

    max_log = np.max(log_w_eff)
    if not np.isfinite(max_log):
        # All weights are zero (or all -inf); fall back to the raw weights.
        return w

    return np.exp(log_w_eff - max_log)


def weighted_pam(
    dist_matrix,
    weights,
    k,
    fitness_scores=None,
    beta=None,
    max_iter=100,
    verbose=False,
):
    n = dist_matrix.shape[1]
    if k <= 0 or k > n:
        raise ValueError("k must be between 1 and N")

    w_eff = effective_weights(weights, fitness_scores, beta)

    remaining = set(range(n))
    medoids = []

    total_weighted_dists = (dist_matrix * w_eff.reshape(1, -1)).sum(axis=1)
    first = int(np.argmin(total_weighted_dists))
    medoids.append(first)
    remaining.remove(first)

    def assign_to_medoids(d_matrix, medoid_list):
        d_to_m = d_matrix[:, medoid_list]
        nearest_idx = np.argmin(d_to_m, axis=1)
        dist1 = d_to_m[np.arange(n), nearest_idx]
        d_to_m_tmp = d_to_m.copy()
        d_to_m_tmp[np.arange(n), nearest_idx] = np.inf
        dist2 = d_to_m_tmp[np.arange(n), np.argmin(d_to_m_tmp, axis=1)]
        nearest_medoid = np.array(medoid_list)[nearest_idx]
        return nearest_idx, nearest_medoid, dist1, dist2

    _, _, dist1, _ = assign_to_medoids(dist_matrix, medoids)

    for _ in range(1, k):
        best_candidate = None
        best_cost = np.inf
        for c in remaining:
            new_dist = np.minimum(dist1, dist_matrix[:, c])
            cost_if_added = float(np.dot(w_eff, new_dist))
            if cost_if_added < best_cost:
                best_cost = cost_if_added
                best_candidate = c
        if best_candidate is None:
            raise RuntimeError("weighted_pam BUILD failed: no valid candidate medoid")
        medoids.append(int(best_candidate))
        remaining.remove(best_candidate)
        _, _, dist1, _ = assign_to_medoids(dist_matrix, medoids)

    if verbose:
        print(f"[BUILD] Initial medoids: {medoids}")

    it = 0
    improved = True
    while improved and it < max_iter:
        improved = False
        best_delta = 0
        best_swap = None
        medoids_set = set(medoids)

        nearest_idx, nearest_medoid, dist1, dist2 = assign_to_medoids(dist_matrix, medoids)

        for m_pos, m in enumerate(medoids):
            for h in range(n):
                if h in medoids_set:
                    continue
                d_ih = dist_matrix[:, h]
                mask_better = d_ih < dist1
                delta_better = ((d_ih - dist1) * w_eff)[mask_better].sum()
                mask_owned_by_m = nearest_medoid == m
                mask_case2 = mask_owned_by_m & (~mask_better)
                new_dist_case2 = np.minimum(dist2, d_ih)
                delta_case2 = ((new_dist_case2 - dist1) * w_eff)[mask_case2].sum()
                delta = delta_better + delta_case2
                if delta < best_delta:
                    best_delta = delta
                    best_swap = (m_pos, h)

        if best_swap and best_delta < -1e-12:
            m_pos, h = best_swap
            medoids[m_pos] = h
            improved = True
        else:
            break
        it += 1

    _, nearest_medoid, dist1, _ = assign_to_medoids(dist_matrix, medoids)
    labels = nearest_medoid.copy()
    return medoids, labels


def future_weighted_pam(
    dist_matrix,
    weights,
    k,
    fitness_scores=None,
    beta=None,
    max_iter=100,
    verbose=False,
):
    """
    Weighted PAM where query points and candidate medoids are different sets.

    dist_matrix shape: (n_points, n_candidates)
    Medoids are indices into the candidate axis (columns).
    """
    n_points, n_candidates = dist_matrix.shape
    if k <= 0 or k > n_candidates:
        raise ValueError("k must be between 1 and the number of candidate medoids")

    w_eff = effective_weights(weights, fitness_scores, beta)

    remaining = set(range(n_candidates))
    medoids = []

    total_weighted_dists = (dist_matrix * w_eff.reshape(-1, 1)).sum(axis=0)
    first = int(np.argmin(total_weighted_dists))
    medoids.append(first)
    remaining.remove(first)

    def assign_to_medoids(d_matrix, medoid_list):
        d_to_m = d_matrix[:, medoid_list]
        nearest_idx = np.argmin(d_to_m, axis=1)
        dist1 = d_to_m[np.arange(n_points), nearest_idx]
        d_to_m_tmp = d_to_m.copy()
        d_to_m_tmp[np.arange(n_points), nearest_idx] = np.inf
        dist2 = d_to_m_tmp[np.arange(n_points), np.argmin(d_to_m_tmp, axis=1)]
        nearest_medoid = np.array(medoid_list)[nearest_idx]
        return nearest_idx, nearest_medoid, dist1, dist2

    _, _, dist1, _ = assign_to_medoids(dist_matrix, medoids)

    for _ in range(1, k):
        best_candidate = None
        best_cost = np.inf
        for c in remaining:
            new_dist = np.minimum(dist1, dist_matrix[:, c])
            cost_if_added = float(np.dot(w_eff, new_dist))
            if cost_if_added < best_cost:
                best_cost = cost_if_added
                best_candidate = c
        if best_candidate is None:
            raise RuntimeError("future_weighted_pam BUILD failed: no valid candidate medoid")
        medoids.append(int(best_candidate))
        remaining.remove(best_candidate)
        _, _, dist1, _ = assign_to_medoids(dist_matrix, medoids)

    if verbose:
        print(f"[BUILD] Initial oracle medoids: {medoids}")

    it = 0
    improved = True
    while improved and it < max_iter:
        improved = False
        best_delta = 0
        best_swap = None
        medoids_set = set(medoids)

        nearest_idx, nearest_medoid, dist1, dist2 = assign_to_medoids(dist_matrix, medoids)

        for m_pos, m in enumerate(medoids):
            for h in range(n_candidates):
                if h in medoids_set:
                    continue
                d_ih = dist_matrix[:, h]
                mask_better = d_ih < dist1
                delta_better = ((d_ih - dist1) * w_eff)[mask_better].sum()
                mask_owned_by_m = nearest_medoid == m
                mask_case2 = mask_owned_by_m & (~mask_better)
                new_dist_case2 = np.minimum(dist2, d_ih)
                delta_case2 = ((new_dist_case2 - dist1) * w_eff)[mask_case2].sum()
                delta = delta_better + delta_case2
                if delta < best_delta:
                    best_delta = delta
                    best_swap = (m_pos, h)

        if best_swap and best_delta < -1e-12:
            m_pos, h = best_swap
            medoids[m_pos] = h
            improved = True
        else:
            break
        it += 1

    _, nearest_medoid, dist1, _ = assign_to_medoids(dist_matrix, medoids)
    labels = nearest_medoid.copy()
    return medoids, labels


def preprocess_for_weighted_pam_fast(
    current_seq_df: pd.DataFrame,
    dist_df: pd.DataFrame,
    *,
    use_covfit_scores: bool = False,
    covfit_score_df: pd.DataFrame = None,
    fitness_column: str = FITNESS_COLUMN,
):
    sequence_ids = current_seq_df["Sequence_ID"].tolist()
    id2idx = {sid: i for i, sid in enumerate(sequence_ids)}
    idx2id = {i: sid for sid, i in id2idx.items()}
    n = len(sequence_ids)

    i_indices = dist_df["Seq_1"].map(id2idx).to_numpy()
    j_indices = dist_df["Seq_2"].map(id2idx).to_numpy()
    distances = dist_df["Distance"].to_numpy(dtype=np.float32)

    dist_matrix = np.zeros((n, n), dtype=np.float32)
    dist_matrix[i_indices, j_indices] = distances

    weights_series = current_seq_df.set_index("Sequence_ID")["Weight"]
    weights = np.array([weights_series[idx2id[i]] for i in range(n)], dtype=np.float32)

    if use_covfit_scores:
        if covfit_score_df is None:
            raise ValueError("covfit_score_df is required when use_covfit_scores=True")
        score_dict = covfit_score_df.set_index("Sequence_ID")[fitness_column].to_dict()
        missing = []
        fitness_values = []
        for i in range(n):
            sid = idx2id[i]
            if sid not in score_dict or pd.isna(score_dict[sid]):
                missing.append(sid)
            else:
                fitness_values.append(score_dict[sid])
        if missing:
            preview = ", ".join(missing[:5])
            suffix = " ..." if len(missing) > 5 else ""
            raise ValueError(
                f"{len(missing)} sequence(s) missing {fitness_column!r}: {preview}{suffix}"
            )
        fitness_scores = np.array(fitness_values, dtype=np.float32)
    else:
        fitness_scores = None

    return dist_matrix, weights, id2idx, idx2id, fitness_scores


def np_hamming_distance_matrix(sequences1, sequences2):
    if not sequences1 or not sequences2:
        return np.array([[]])
    arr1 = np.array([list(s) for s in sequences1])
    arr2 = np.array([list(s) for s in sequences2])
    return np.sum(np.not_equal(arr1[:, None, :], arr2[None, :, :]), axis=2)


def _hamming_distance_matrix_local(sequences1, sequences2, max_block_bytes=200_000_000):
    if not sequences1 or not sequences2:
        return np.zeros((len(sequences1), len(sequences2)), dtype=np.float32)

    n1, n2 = len(sequences1), len(sequences2)
    seq_len = len(sequences1[0])
    out = np.zeros((n1, n2), dtype=np.float32)
    cand = np.array([list(s) for s in sequences2])

    elems_per_block = max(1, max_block_bytes // max(seq_len, 1))
    point_batch = max(1, min(n1, int(np.sqrt(elems_per_block))))
    cand_batch = max(1, min(n2, elems_per_block // point_batch))

    for c0 in range(0, n2, cand_batch):
        c1 = min(c0 + cand_batch, n2)
        cand_chunk = cand[c0:c1]
        for p0 in range(0, n1, point_batch):
            p1 = min(p0 + point_batch, n1)
            point_chunk = np.array([list(s) for s in sequences1[p0:p1]])
            out[p0:p1, c0:c1] = np.sum(
                point_chunk[:, None, :] != cand_chunk[None, :, :],
                axis=2,
            )
    return out


def fitness_adjusted_weighted_min_distance(
    distance_matrix,
    seq_fitness,
    candidate_fitness,
    weights,
    fitness_weighted=False,
    beta=None,
):
    """
    Weighted average of minimum Hamming distances.

    When fitness_weighted is False:
        contrib_i = w_i * min_j(distance_ij)

    When fitness_weighted is True (current circulating pool):
        contrib_i = w_i * e ** (beta * seq_i) * min_j(distance_ij / e ** (beta * candidate_j))
    """
    weights = np.asarray(weights, dtype=float)

    if distance_matrix.size == 0:
        return 0.0

    if fitness_weighted:
        if beta is None:
            raise ValueError("beta is required when fitness_weighted=True")
        seq_fitness_modifiers = exp_fitness_modifier(seq_fitness, beta).astype(np.float64)
        candidate_fitness_modifiers = exp_fitness_modifier(
            candidate_fitness, beta
        ).astype(np.float64)
        distance_matrix = distance_matrix / candidate_fitness_modifiers.reshape(1, -1)
        weights = weights * seq_fitness_modifiers
    else:
        del seq_fitness, candidate_fitness

    min_distances = np.min(distance_matrix, axis=1)

    total_weight = np.sum(weights)
    if total_weight == 0:
        return 0.0
    return float(np.sum(min_distances * weights) / total_weight)


def evaluate_medoids_vectorized(medoids, id_to_seq_, future_seq_dfs, current_seq_df, fitness_column=FITNESS_COLUMN):
    medoid_seqs = [id_to_seq_[mid] for mid in medoids]
    medoid_fitness = current_seq_df.set_index("Sequence_ID").loc[medoids, fitness_column].values
    results = {}

    for interval, df in future_seq_dfs.items():
        if df.empty:
            results[interval] = 0.0
            continue

        future_seq_ids = df["Sequence_ID"].tolist()
        future_seqs = [id_to_seq_[sid] for sid in future_seq_ids]
        weights = df["Weight"].values
        seq_fitness = df[fitness_column].values

        if not future_seqs:
            results[interval] = 0.0
            continue

        distance_matrix = np_hamming_distance_matrix(future_seqs, medoid_seqs)
        results[interval] = fitness_adjusted_weighted_min_distance(
            distance_matrix=distance_matrix,
            seq_fitness=seq_fitness,
            candidate_fitness=medoid_fitness,
            weights=weights,
        )
    return results


def evaluate_on_current_vectorized(
    candidate_seqs,
    current_seq_df,
    id_to_seq,
    beta,
    fitness_column=FITNESS_COLUMN,
):
    """
    Fitness-weighted average minimum Hamming distance from circulating sequences
    to the nearest candidate sequence.

    contrib_i = w_i * e ** (beta * fitness_i) * min_j(distance_ij)

    The effective weights come from the same `effective_weights` helper the PAM
    medoid search uses, so the weighting here is identical to the one driving
    medoid selection: built in log-space and rescaled by the max so the
    e ** (beta * fitness) term cannot overflow.
    """
    if not candidate_seqs or current_seq_df.empty:
        return 0.0

    current_ids = current_seq_df["Sequence_ID"].tolist()
    current_seqs = [id_to_seq[sid] for sid in current_ids]
    weights = np.asarray(current_seq_df["Weight"].values, dtype=float)
    seq_fitness = np.asarray(current_seq_df[fitness_column].values, dtype=float)
    weight_sum = float(np.sum(weights))
    if weight_sum == 0:
        return 0.0

    distance_matrix = np_hamming_distance_matrix(current_seqs, candidate_seqs)
    min_distances = np.min(distance_matrix, axis=1).astype(np.float64)
    w_eff = effective_weights(weights, seq_fitness, beta)

    return float(np.sum(min_distances * w_eff) / float(np.sum(w_eff)))


def evaluate_vaccines_vectorized(
    vaccines,
    vaccine_to_seq,
    id_to_seq,
    future_seq_dfs,
    vaccine_fitness,
    fitness_column=FITNESS_COLUMN,
):
    vaccine_ids = [vid for vid in vaccines if vid in vaccine_to_seq]
    candidate_seqs = [vaccine_to_seq[vid] for vid in vaccine_ids]
    candidate_fitness = np.array([vaccine_fitness[vid] for vid in vaccine_ids], dtype=float)

    if not candidate_seqs:
        return {interval: 0.0 for interval in future_seq_dfs}

    results = {}
    for interval, df in future_seq_dfs.items():
        if df.empty:
            results[interval] = 0.0
            continue

        valid_df = df[df["Sequence_ID"].isin(id_to_seq)]
        if valid_df.empty:
            results[interval] = 0.0
            continue

        future_seq_ids = valid_df["Sequence_ID"].tolist()
        future_seqs = [id_to_seq[sid] for sid in future_seq_ids]
        weights = valid_df["Weight"].values
        seq_fitness = valid_df[fitness_column].values

        distance_matrix = np_hamming_distance_matrix(future_seqs, candidate_seqs)
        results[interval] = fitness_adjusted_weighted_min_distance(
            distance_matrix=distance_matrix,
            seq_fitness=seq_fitness,
            candidate_fitness=candidate_fitness,
            weights=weights,
        )
    return results


def find_pairwise_dist_bin() -> Path:
    for base in [Path.cwd().resolve(), REPO_ROOT, SCRIPT_DIR]:
        for _ in range(6):
            candidate = base / "scripts" / "pairwise_dist"
            if candidate.is_file():
                return candidate
            if base.parent == base:
                break
            base = base.parent
    raise FileNotFoundError(
        "pairwise_dist not found. Build with: "
        "cmake -S scripts -B scripts/build && cmake --build scripts/build -j"
    )


def find_covfit_dir() -> Path:
    covfit_dir = REPO_ROOT / "CoVFit_CLI"
    if not (covfit_dir / "covfit_cli").exists():
        raise FileNotFoundError(f"covfit_cli not found in {covfit_dir}")
    return covfit_dir


def run_covfit(fasta_path: Path, outdir: Path, covfit_dir: Path) -> Path:
    outdir.mkdir(parents=True, exist_ok=True)
    output_tsv = outdir / "CoVFit_Predictions_fold_3.tsv"
    if output_tsv.exists():
        print(f"  Reusing CoVFit output: {output_tsv}")
        return output_tsv

    cmd = [
        "./covfit_cli",
        "-i",
        str(fasta_path),
        "-o",
        str(outdir),
        "-g",
        "-b",
        "16",
        "-f",
        "3",
        "-d",
    ]
    print(f"  Running CoVFit: {' '.join(cmd)}")
    result = subprocess.run(cmd, cwd=str(covfit_dir), capture_output=True, text=True)
    if result.stdout:
        print(result.stdout)
    if result.stderr:
        print(result.stderr, file=sys.stderr)
    if result.returncode != 0:
        raise RuntimeError(f"covfit_cli failed with exit code {result.returncode}")
    return output_tsv


def load_covfit_scores(tsv_path: Path) -> pd.DataFrame:
    df = pd.read_csv(tsv_path, sep="\t")
    return df.rename(columns={"seq_name": "Sequence_ID"})


def attach_fitness_column(df: pd.DataFrame, score_map: pd.Series, label: str) -> pd.DataFrame:
    missing = set(df["Sequence_ID"]) - set(score_map.index)
    if missing:
        raise ValueError(f"{len(missing)} sequence(s) missing {label} in {sorted(missing)[:5]}")
    return df.assign(**{FITNESS_COLUMN: df["Sequence_ID"].map(score_map)})


def run_pairwise_dist(current_seq_df: pd.DataFrame, id_to_seq: dict, cohort_dir: Path, skip: bool) -> pd.DataFrame:
    cohort_dir.mkdir(parents=True, exist_ok=True)
    current_dict = cohort_dir / "current_id_to_seq.parquet"
    dist_parquet = cohort_dir / "pairwise_dist.parquet"

    if skip and dist_parquet.exists():
        print(f"  Reusing pairwise distances: {dist_parquet}")
        return pd.read_parquet(dist_parquet, engine="pyarrow")

    current_seq_ids = current_seq_df["Sequence_ID"].tolist()
    pq.write_table(
        pa.table(
            {
                "key": pa.array(current_seq_ids, type=pa.string()),
                "value": pa.array([id_to_seq[sid] for sid in current_seq_ids], type=pa.string()),
            }
        ),
        str(current_dict),
    )

    pairwise_dist_bin = find_pairwise_dist_bin()
    cmd = [
        str(pairwise_dist_bin.resolve()),
        "--parquet",
        str(current_dict),
        "--output",
        str(dist_parquet),
    ]
    print(f"  Running pairwise_dist: {' '.join(cmd)}")
    subprocess.run(cmd, check=True)
    return pd.read_parquet(dist_parquet, engine="pyarrow")


def build_cohort_data(final_df: pd.DataFrame, cutoff_date: str) -> Tuple[pd.DataFrame, Dict[str, pd.DataFrame]]:
    current_df, future_df = split_circulating_and_future(
        final_df,
        cutoff_date_str=cutoff_date,
        left_buffer_days=4,
        right_buffer_days=4,
        circulation_months=None,
        circulation_days=CIRCULATION_DAYS,
        future_months=12,
        future_days=None,
    )
    future_slices = slice_by_disjoint_windows(future_df, DISJOINT_WINDOWS)

    current_seq_df = (
        current_df.groupby("Sequence_ID")
        .size()
        .reset_index(name="Tip_Count")
        .sort_values(by="Tip_Count", ascending=False)
        .reset_index(drop=True)
    )
    current_seq_df = current_seq_df[current_seq_df["Tip_Count"] >= MIN_TIP_COUNT].reset_index(drop=True)
    current_seq_df["Weight"] = current_seq_df["Tip_Count"] / current_seq_df["Tip_Count"].sum()

    future_seq_dfs = {}
    for key, frame in future_slices.items():
        grouped = (
            frame.groupby("Sequence_ID")
            .size()
            .reset_index(name="Tip_Count")
            .sort_values(by="Tip_Count", ascending=False)
            .reset_index(drop=True)
        )
        grouped["Weight"] = grouped["Tip_Count"] / grouped["Tip_Count"].sum()
        future_seq_dfs[key] = grouped

    return current_seq_df, future_seq_dfs


def beta_log_suffix(beta: float) -> str:
    """Stable beta token for lineage log filenames (60.0 -> '60')."""
    return f"{beta:g}"


def medoid_lineage_log_path(root_path: Path, cutoff_date: str, beta: float) -> Path:
    return (
        root_path
        / cutoff_date
        / f"medoid_lineage_assignments_linear_exp_{beta_log_suffix(beta)}.log"
    )


def legacy_medoid_lineage_log_path(root_path: Path, beta: float) -> Path:
    return root_path / f"medoid_lineage_assignments_linear_exp_{beta_log_suffix(beta)}.log"


def spike_mutation_distance(seq_a: str, seq_b: str) -> int:
    if len(seq_a) != len(seq_b):
        raise ValueError(f"Sequence length mismatch: {len(seq_a)} vs {len(seq_b)}")
    return sum(a != b for a, b in zip(seq_a, seq_b))


def load_lineage_spike_groups(fasta_path: Path):
    groups = []
    for record in SeqIO.parse(str(fasta_path), "fasta"):
        groups.append(
            {
                "header": record.id,
                "lineages": record.id.split("|"),
                "sequence": str(record.seq),
            }
        )
    return groups


def assign_medoid_lineages(medoid_seq: str, lineage_groups):
    exact = [g for g in lineage_groups if medoid_seq == g["sequence"]]
    if exact:
        return {
            "match_type": "exact",
            "lineage_groups": [g["header"] for g in exact],
            "distance": 0,
            "closest": [(g["header"], 0) for g in exact],
        }

    distances = [
        (g["header"], spike_mutation_distance(medoid_seq, g["sequence"]))
        for g in lineage_groups
    ]
    min_dist = min(d for _, d in distances)
    closest = [(header, dist) for header, dist in distances if dist == min_dist]
    closest.sort(key=lambda x: x[0])
    return {
        "match_type": "closest",
        "lineage_groups": [header for header, _ in closest],
        "distance": min_dist,
        "closest": closest,
    }


def first_lineage_from_group(lineage_group: str) -> str:
    return lineage_group.split("|")[0].strip()


def write_medoid_lineage_log(
    pam_results: Dict[int, dict],
    id_to_seq: dict,
    lineage_groups,
    log_path: Path,
    lineage_fasta: Path,
) -> Dict[Tuple[int, str], str]:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    lineage_map: Dict[Tuple[int, str], str] = {}
    log_lines = [
        "Medoid lineage assignment log",
        f"Generated: {datetime.now().isoformat(timespec='seconds')}",
        f"Lineage reference: {lineage_fasta}",
        f"N lineage spike groups: {len(lineage_groups)}",
        "",
    ]

    for k, res in sorted(pam_results.items()):
        log_lines.append(f"=== K = {k} ===")
        for sid in res["medoids"]:
            seq = id_to_seq.get(sid)
            log_lines.append(f"K={k}\tSequence_ID={sid}")
            if seq is None:
                log_lines.append("  status: missing from id_to_seq")
                lineage_map[(k, sid)] = ""
                continue

            assignment = assign_medoid_lineages(seq, lineage_groups)
            log_lines.append(f"  match_type: {assignment['match_type']}")
            log_lines.append(f"  mutation_distance: {assignment['distance']}")
            if assignment["match_type"] == "exact":
                header = assignment["lineage_groups"][0]
                log_lines.append(f"  lineage_group: {header}")
                lineage_map[(k, sid)] = first_lineage_from_group(header)
            else:
                log_lines.append("  closest_lineage_groups:")
                for header, dist in assignment["closest"]:
                    log_lines.append(f"    - {header} ({dist})")
                lineage_map[(k, sid)] = first_lineage_from_group(assignment["lineage_groups"][0])

        log_lines.append("")

    log_path.write_text("\n".join(log_lines) + "\n")
    print(f"  Wrote lineage log: {log_path}")
    return lineage_map


def parse_medoid_lineage_log(log_path: Path) -> Dict[Tuple[int, str], str]:
    if not log_path.exists():
        raise FileNotFoundError(f"Lineage log not found: {log_path}")

    text = log_path.read_text()
    lineage_map: Dict[Tuple[int, str], str] = {}
    current_k = None
    current_sid = None

    for line in text.splitlines():
        header_match = re.match(r"^=== K = (\d+) ===$", line.strip())
        if header_match:
            current_k = int(header_match.group(1))
            continue

        medoid_match = re.match(r"^K=(\d+)\tSequence_ID=(\S+)$", line.strip())
        if medoid_match:
            current_k = int(medoid_match.group(1))
            current_sid = medoid_match.group(2)
            continue

        lineage_match = re.match(r"^\s*lineage_group:\s*(.+)$", line)
        if lineage_match and current_k is not None and current_sid is not None:
            lineage_map[(current_k, current_sid)] = first_lineage_from_group(lineage_match.group(1).strip())
            current_sid = None

    return lineage_map


def cutoff_display_name(cutoff_date: str) -> str:
    year, month, _ = cutoff_date.split("-")
    month_names = {
        "01": "Jan", "02": "Feb", "03": "Mar", "04": "Apr", "05": "May", "06": "Jun",
        "07": "Jul", "08": "Aug", "09": "Sep", "10": "Oct", "11": "Nov", "12": "Dec",
    }
    return f"{month_names[month]}-{year[2:]}"


def modifier_display(beta: float) -> str:
    return f"e ^ ({beta} * fitness)"


def pyro_fitness_for_medoids(medoid_ids: Iterable[str], id_to_seq: dict, ref_seq: str, mutation_score_dict: dict) -> Dict[str, float]:
    scores = {}
    for sid in medoid_ids:
        seq = id_to_seq.get(sid)
        if seq is None:
            scores[sid] = np.nan
        else:
            mutations = compute_mutations(seq, ref_seq)
            scores[sid] = compute_fitness_score_pyro(mutations, mutation_score_dict)
    return scores


def prepare_cutoff_assets(
    root_path: Path,
    cutoff_date: str,
    final_df: pd.DataFrame,
    id_to_seq: dict,
    covfit_dir: Path,
    skip_pairwise: bool,
) -> Tuple[pd.DataFrame, Dict[str, pd.DataFrame], np.ndarray, np.ndarray, dict, dict, np.ndarray, pd.DataFrame]:
    print(f"\n=== Preparing cohort for cutoff {cutoff_date} ===")
    cohort_dir = root_path / cutoff_date
    current_seq_df, future_seq_dfs = build_cohort_data(final_df, cutoff_date)
    print(f"  Current sequences: {len(current_seq_df):,}")

    current_fasta = cohort_dir / "vaccine_scout_current_seq_df.fasta"
    with open(current_fasta, "w") as handle:
        for seq_id in current_seq_df["Sequence_ID"]:
            handle.write(f">{seq_id}\n{id_to_seq[seq_id]}\n")

    current_covfit_tsv = run_covfit(
        current_fasta,
        cohort_dir / "current_covfit",
        covfit_dir,
    )

    covfit_score_df = load_covfit_scores(current_covfit_tsv)
    score_map = covfit_score_df.set_index("Sequence_ID")[FITNESS_COLUMN]
    current_seq_df = attach_fitness_column(current_seq_df, score_map, "current CoVFit")

    future_ids = (
        pd.concat([df[["Sequence_ID"]] for df in future_seq_dfs.values()])
        .drop_duplicates()["Sequence_ID"]
        .tolist()
    )
    future_fasta = cohort_dir / "vaccine_scout_future_seq_df.fasta"
    with open(future_fasta, "w") as handle:
        for seq_id in future_ids:
            handle.write(f">{seq_id}\n{id_to_seq[seq_id]}\n")

    future_covfit_tsv = run_covfit(
        future_fasta,
        cohort_dir / "future_covfit",
        covfit_dir,
    )

    future_covfit_df = load_covfit_scores(future_covfit_tsv)
    future_score_map = future_covfit_df.set_index("Sequence_ID")[FITNESS_COLUMN]
    for interval in future_seq_dfs:
        future_seq_dfs[interval] = attach_fitness_column(
            future_seq_dfs[interval], future_score_map, f"future {interval}"
        )

    dist_df = run_pairwise_dist(current_seq_df, id_to_seq, cohort_dir, skip_pairwise)
    dist_matrix, weights, id2idx, idx2id, fitness_scores = preprocess_for_weighted_pam_fast(
        current_seq_df,
        dist_df,
        use_covfit_scores=True,
        covfit_score_df=covfit_score_df,
    )

    return current_seq_df, future_seq_dfs, dist_matrix, weights, id2idx, idx2id, fitness_scores, dist_df


def run_beta_analysis(
    beta: float,
    dist_matrix: np.ndarray,
    weights: np.ndarray,
    idx2id: dict,
    fitness_scores: np.ndarray,
    current_seq_df: pd.DataFrame,
    future_seq_dfs: Dict[str, pd.DataFrame],
    id_to_seq: dict,
    ref_seq: str,
    mutation_score_dict: dict,
    lineage_groups,
    lineage_fasta: Path,
    root_path: Path,
    cutoff_date: str,
    write_lineage: bool,
) -> List[dict]:
    print(f"\n--- BETA = {beta} ---")

    pam_results = {}
    for k in K_VALUES:
        medoids, labels = weighted_pam(
            dist_matrix,
            weights,
            k=k,
            fitness_scores=fitness_scores,
            beta=beta,
            max_iter=100,
            verbose=False,
        )
        pam_results[k] = {
            "medoids": [idx2id[i] for i in medoids],
            "labels": [idx2id[i] for i in labels],
        }

    log_path = medoid_lineage_log_path(root_path, cutoff_date, beta)
    if write_lineage:
        lineage_map = write_medoid_lineage_log(
            pam_results, id_to_seq, lineage_groups, log_path, lineage_fasta
        )
    else:
        fallback = legacy_medoid_lineage_log_path(root_path, beta)
        lineage_map = parse_medoid_lineage_log(log_path if log_path.exists() else fallback)

    covfit_map = current_seq_df.set_index("Sequence_ID")[FITNESS_COLUMN].to_dict()

    rows = []
    for k in K_VALUES:
        medoid_ids = pam_results[k]["medoids"]
        pyro_map = pyro_fitness_for_medoids(medoid_ids, id_to_seq, ref_seq, mutation_score_dict)

        eval_future = evaluate_medoids_vectorized(
            medoids=medoid_ids,
            id_to_seq_=id_to_seq,
            future_seq_dfs=future_seq_dfs,
            current_seq_df=current_seq_df,
        )
        medoid_seqs = [id_to_seq[mid] for mid in medoid_ids]
        dist_current = evaluate_on_current_vectorized(
            candidate_seqs=medoid_seqs,
            current_seq_df=current_seq_df,
            id_to_seq=id_to_seq,
            beta=CURRENT_DISTANCE_BETA,
        )

        for sid in medoid_ids:
            rows.append(
                {
                    "cutoff_date": cutoff_date,
                    "beta": beta,
                    "k": k,
                    "weights_function": "Linear",
                    "neutralization_function": modifier_display(beta),
                    "medoid": sid,
                    "lineage": lineage_map.get((k, sid), ""),
                    "covfit_fitness": covfit_map.get(sid, np.nan),
                    "pyro_fitness": pyro_map.get(sid, np.nan),
                    "distance_current": dist_current,
                    "distance_0_6": eval_future.get("6M_1", np.nan),
                    "distance_6_12": eval_future.get("6M_2", np.nan),
                }
            )

    return rows


def prepare_vaccine_covfit(
    root_path: Path,
    cutoff_date: str,
    covfit_dir: Path,
) -> Tuple[dict, Dict[str, float]]:
    vaccine_df = read_pickle(root_path / "vaccine_recommendations.pkl")
    vaccine_df = vaccine_df.rename(columns={"clade": "Sequence_ID"})
    clade_to_seq = dict(zip(vaccine_df["Sequence_ID"], vaccine_df["mutated_sequence"]))

    cohort_dir = root_path / cutoff_date
    vaccine_fasta = cohort_dir / "vaccine_scout_strains.fasta"
    with open(vaccine_fasta, "w") as handle:
        for strain, seq in clade_to_seq.items():
            handle.write(f">{strain}\n{seq}\n")

    vaccine_covfit_tsv = run_covfit(
        vaccine_fasta,
        cohort_dir / "vaccine_covfit",
        covfit_dir,
    )

    vaccine_covfit_df = load_covfit_scores(vaccine_covfit_tsv)
    vaccine_fitness = vaccine_covfit_df.set_index("Sequence_ID")[FITNESS_COLUMN].to_dict()
    return clade_to_seq, vaccine_fitness


def load_pfizer_sequence(cutoff_date: str) -> Tuple[str, str]:
    if cutoff_date not in PFIZER_INN_BY_CUTOFF:
        raise ValueError(f"No Pfizer vaccine mapping for cutoff_date={cutoff_date!r}")

    pfizer_inn = PFIZER_INN_BY_CUTOFF[cutoff_date]
    fasta_name = PFIZER_FASTA_BY_CUTOFF[cutoff_date]
    fasta_path = REPO_ROOT / "data" / "Pfizer" / fasta_name
    if not fasta_path.is_file():
        raise FileNotFoundError(f"Pfizer AA FASTA not found: {fasta_path}")

    record = next(SeqIO.parse(str(fasta_path), "fasta"))
    return pfizer_inn, str(record.seq)


def prepare_pfizer_covfit(
    root_path: Path,
    cutoff_date: str,
    covfit_dir: Path,
    pfizer_inn: str,
    pfizer_seq: str,
) -> Dict[str, float]:
    cohort_dir = root_path / cutoff_date
    pfizer_fasta = cohort_dir / f"pfizer_{pfizer_inn}.fasta"
    with open(pfizer_fasta, "w") as handle:
        handle.write(f">{pfizer_inn}\n{pfizer_seq}\n")

    pfizer_covfit_tsv = run_covfit(
        pfizer_fasta,
        cohort_dir / "pfizer_covfit",
        covfit_dir,
    )
    pfizer_covfit_df = load_covfit_scores(pfizer_covfit_tsv)
    pfizer_fitness = pfizer_covfit_df.set_index("Sequence_ID")[FITNESS_COLUMN].to_dict()
    if pfizer_inn not in pfizer_fitness:
        raise ValueError(
            f"CoVFit output missing Pfizer strain {pfizer_inn!r}; "
            f"available: {list(pfizer_fitness.keys())}"
        )
    return pfizer_fitness


def run_fda_variant_analysis(
    strain: str,
    clade_to_seq: dict,
    vaccine_fitness: Dict[str, float],
    current_seq_df: pd.DataFrame,
    future_seq_dfs: Dict[str, pd.DataFrame],
    id_to_seq: dict,
    ref_seq: str,
    mutation_score_dict: dict,
    lineage_groups,
) -> dict:
    if strain not in clade_to_seq:
        raise ValueError(f"FDA strain {strain!r} not found in vaccine recommendations")

    vaccine_seq = clade_to_seq[strain]
    assignment = assign_medoid_lineages(vaccine_seq, lineage_groups)
    if assignment["match_type"] == "exact":
        lineage = first_lineage_from_group(assignment["lineage_groups"][0])
    else:
        lineage = first_lineage_from_group(assignment["lineage_groups"][0])

    mutations = compute_mutations(vaccine_seq, ref_seq)
    pyro_fitness = compute_fitness_score_pyro(mutations, mutation_score_dict)
    covfit_fitness = vaccine_fitness[strain]

    eval_future = evaluate_vaccines_vectorized(
        vaccines=[strain],
        vaccine_to_seq=clade_to_seq,
        id_to_seq=id_to_seq,
        future_seq_dfs=future_seq_dfs,
        vaccine_fitness=vaccine_fitness,
    )
    dist_current = evaluate_on_current_vectorized(
        candidate_seqs=[vaccine_seq],
        current_seq_df=current_seq_df,
        id_to_seq=id_to_seq,
        beta=CURRENT_DISTANCE_BETA,
    )

    return {
        "strain": strain,
        "lineage": lineage,
        "covfit_fitness": covfit_fitness,
        "pyro_fitness": pyro_fitness,
        "distance_current": dist_current,
        "distance_0_6": eval_future.get("6M_1", np.nan),
        "distance_6_12": eval_future.get("6M_2", np.nan),
    }


def run_pfizer_variant_analysis(
    pfizer_inn: str,
    pfizer_seq: str,
    pfizer_fitness: Dict[str, float],
    current_seq_df: pd.DataFrame,
    future_seq_dfs: Dict[str, pd.DataFrame],
    id_to_seq: dict,
    ref_seq: str,
    mutation_score_dict: dict,
    lineage_groups,
) -> dict:
    pfizer_to_seq = {pfizer_inn: pfizer_seq}
    assignment = assign_medoid_lineages(pfizer_seq, lineage_groups)
    lineage = first_lineage_from_group(assignment["lineage_groups"][0])

    mutations = compute_mutations(pfizer_seq, ref_seq)
    pyro_fitness = compute_fitness_score_pyro(mutations, mutation_score_dict)
    covfit_fitness = pfizer_fitness[pfizer_inn]

    eval_future = evaluate_vaccines_vectorized(
        vaccines=[pfizer_inn],
        vaccine_to_seq=pfizer_to_seq,
        id_to_seq=id_to_seq,
        future_seq_dfs=future_seq_dfs,
        vaccine_fitness=pfizer_fitness,
    )
    dist_current = evaluate_on_current_vectorized(
        candidate_seqs=[pfizer_seq],
        current_seq_df=current_seq_df,
        id_to_seq=id_to_seq,
        beta=CURRENT_DISTANCE_BETA,
    )

    return {
        "strain": pfizer_inn,
        "lineage": lineage,
        "covfit_fitness": covfit_fitness,
        "pyro_fitness": pyro_fitness,
        "distance_current": dist_current,
        "distance_0_6": eval_future.get("6M_1", np.nan),
        "distance_6_12": eval_future.get("6M_2", np.nan),
    }


def build_oracle_distance_matrix(
    oracle_future_df: pd.DataFrame,
    candidate_ids: List[str],
    point_ids: List[str],
    dist_df: pd.DataFrame,
    id_to_seq: dict,
) -> np.ndarray:
    point_id2row = {sid: i for i, sid in enumerate(point_ids)}
    cand_id2idx = {sid: i for i, sid in enumerate(candidate_ids)}
    in_current = oracle_future_df["Sequence_ID"].isin(candidate_ids)

    n_points = len(point_ids)
    n_candidates = len(candidate_ids)
    oracle_dist_matrix = np.full((n_points, n_candidates), np.nan, dtype=np.float32)

    for seq_col, other_col in [("Seq_1", "Seq_2"), ("Seq_2", "Seq_1")]:
        sub = dist_df[
            dist_df[seq_col].isin(candidate_ids)
            & dist_df[other_col].isin(candidate_ids)
        ]
        for row in sub.itertuples(index=False):
            point_id = getattr(row, seq_col)
            cand_id = getattr(row, other_col)
            i = point_id2row.get(point_id)
            j = cand_id2idx.get(cand_id)
            if i is not None and j is not None:
                oracle_dist_matrix[i, j] = row.Distance

    future_only_rows = np.flatnonzero(~in_current.to_numpy())
    if len(future_only_rows) > 0:
        future_only_seqs = [id_to_seq[point_ids[i]] for i in future_only_rows]
        candidate_seqs = [id_to_seq[sid] for sid in candidate_ids]
        hamming_block = _hamming_distance_matrix_local(
            future_only_seqs, candidate_seqs
        )
        oracle_dist_matrix[np.array(future_only_rows), :] = hamming_block

    missing_pairs = np.isnan(oracle_dist_matrix).sum()
    if missing_pairs:
        raise ValueError(
            f"Oracle distance matrix has {missing_pairs:,} missing entries; "
            "check pairwise_dist coverage or id_to_seq for future-only sequences."
        )
    return oracle_dist_matrix


def run_oracle_analysis(
    cutoff_date: str,
    current_seq_df: pd.DataFrame,
    future_seq_dfs: Dict[str, pd.DataFrame],
    dist_df: pd.DataFrame,
    id_to_seq: dict,
    ref_seq: str,
    mutation_score_dict: dict,
    lineage_groups,
) -> Dict[str, Dict[int, List[dict]]]:
    """
    Run oracle PAM for each future slice (6M_1, 6M_2).

    Returns:
        {interval: {k: [per-medoid row dicts]}}
    """
    current_ids = current_seq_df["Sequence_ID"]
    covfit_map = current_seq_df.set_index("Sequence_ID")[FITNESS_COLUMN].to_dict()
    results: Dict[str, Dict[int, List[dict]]] = {}

    for oracle_interval in DISJOINT_WINDOWS:
        print(f"\n--- Oracle PAM ({oracle_interval}) ---")
        oracle_future_df = future_seq_dfs[oracle_interval].copy().reset_index(drop=True)
        point_ids = oracle_future_df["Sequence_ID"].tolist()
        in_current = oracle_future_df["Sequence_ID"].isin(current_ids)
        candidate_ids = oracle_future_df.loc[in_current, "Sequence_ID"].tolist()

        if not candidate_ids:
            raise ValueError(
                f"No overlapping sequences between {oracle_interval} and current_seq_df"
            )

        print(
            f"  {len(point_ids):,} future points, "
            f"{len(candidate_ids):,} circulating candidates"
        )

        oracle_dist_matrix = build_oracle_distance_matrix(
            oracle_future_df=oracle_future_df,
            candidate_ids=candidate_ids,
            point_ids=point_ids,
            dist_df=dist_df,
            id_to_seq=id_to_seq,
        )
        oracle_point_weights = oracle_future_df["Weight"].to_numpy(dtype=np.float32)
        oracle_point_fitness = oracle_future_df[FITNESS_COLUMN].to_numpy(dtype=np.float32)

        interval_results: Dict[int, List[dict]] = {}
        for k in K_VALUES:
            medoid_idx, _labels = future_weighted_pam(
                oracle_dist_matrix,
                oracle_point_weights,
                k=k,
                fitness_scores=oracle_point_fitness,
                beta=0,
                max_iter=100,
                verbose=False,
            )
            medoid_ids = [candidate_ids[i] for i in medoid_idx]
            print(f"  K={k}: medoids={medoid_ids}")

            pyro_map = pyro_fitness_for_medoids(
                medoid_ids, id_to_seq, ref_seq, mutation_score_dict
            )
            eval_future = evaluate_medoids_vectorized(
                medoids=medoid_ids,
                id_to_seq_=id_to_seq,
                future_seq_dfs=future_seq_dfs,
                current_seq_df=current_seq_df,
            )
            medoid_seqs = [id_to_seq[sid] for sid in medoid_ids]
            dist_current = evaluate_on_current_vectorized(
                candidate_seqs=medoid_seqs,
                current_seq_df=current_seq_df,
                id_to_seq=id_to_seq,
                beta=CURRENT_DISTANCE_BETA,
            )

            k_rows = []
            for sid in medoid_ids:
                seq = id_to_seq.get(sid)
                if seq is None:
                    lineage = ""
                else:
                    assignment = assign_medoid_lineages(seq, lineage_groups)
                    lineage = first_lineage_from_group(assignment["lineage_groups"][0])

                k_rows.append(
                    {
                        "cutoff_date": cutoff_date,
                        "oracle_interval": oracle_interval,
                        "fit_slice_label": ORACLE_INTERVAL_LABELS[oracle_interval],
                        "k": k,
                        "medoid": sid,
                        "lineage": lineage,
                        "covfit_fitness": covfit_map.get(sid, np.nan),
                        "pyro_fitness": pyro_map.get(sid, np.nan),
                        "distance_current": dist_current,
                        "distance_0_6": eval_future.get("6M_1", np.nan),
                        "distance_6_12": eval_future.get("6M_2", np.nan),
                    }
                )
            interval_results[k] = k_rows

        results[oracle_interval] = interval_results

    return results


def nest_medoid_rows(all_rows: List[dict]) -> Dict[str, Dict[int, Dict[float, List[dict]]]]:
    nested: Dict[str, Dict[int, Dict[float, List[dict]]]] = {}
    for row in all_rows:
        nested.setdefault(row["cutoff_date"], {}).setdefault(row["k"], {}).setdefault(row["beta"], []).append(row)
    return nested


def sheet_total_columns(betas: List[float]) -> int:
    return len(betas) * BLOCK_COLS + max(0, len(betas) - 1) * BLOCK_GUTTER


def block_start_column(beta_idx: int) -> int:
    return 1 + beta_idx * (BLOCK_COLS + BLOCK_GUTTER)


def write_beta_header_row(ws, row: int, betas: List[float], header_font: Font):
    for beta_idx, beta in enumerate(betas):
        start_col = block_start_column(beta_idx)
        for offset, header in enumerate(EXCEL_COLUMNS):
            cell = ws.cell(row=row, column=start_col + offset, value=header)
            cell.font = header_font


def write_medoid_block(
    ws,
    first_row: int,
    last_row: int,
    start_col: int,
    medoid_rows: List[dict],
    beta: float,
):
    if not medoid_rows:
        return

    shared = medoid_rows[0]
    merge_cols = [0, 1, 6, 7, 8]

    for col_offset in merge_cols:
        col = start_col + col_offset
        if last_row > first_row:
            ws.merge_cells(
                start_row=first_row,
                start_column=col,
                end_row=last_row,
                end_column=col,
            )
        value = {
            0: shared["weights_function"],
            1: modifier_display(beta),
            6: shared["distance_current"],
            7: shared["distance_0_6"],
            8: shared["distance_6_12"],
        }[col_offset]
        cell = ws.cell(row=first_row, column=col, value=value)
        cell.alignment = Alignment(vertical="center")

    for row_offset, medoid_row in enumerate(medoid_rows):
        row = first_row + row_offset
        ws.cell(row=row, column=start_col + 2, value=medoid_row["medoid"])
        ws.cell(row=row, column=start_col + 3, value=medoid_row["lineage"])
        ws.cell(row=row, column=start_col + 4, value=medoid_row["covfit_fitness"])
        ws.cell(row=row, column=start_col + 5, value=medoid_row["pyro_fitness"])


def write_reference_vaccine_header_row(
    ws,
    row: int,
    header_font: Font,
    columns: List[str],
):
    start_col = block_start_column(0)
    for offset, header in enumerate(columns):
        cell = ws.cell(row=row, column=start_col + offset, value=header)
        cell.font = header_font


def write_reference_vaccine_row(ws, row: int, vaccine_data: dict):
    start_col = block_start_column(0)
    ws.cell(row=row, column=start_col + 0, value=vaccine_data["strain"])
    ws.cell(row=row, column=start_col + 3, value=vaccine_data["lineage"])
    ws.cell(row=row, column=start_col + 4, value=vaccine_data["covfit_fitness"])
    ws.cell(row=row, column=start_col + 5, value=vaccine_data["pyro_fitness"])
    ws.cell(row=row, column=start_col + 6, value=vaccine_data["distance_current"])
    ws.cell(row=row, column=start_col + 7, value=vaccine_data["distance_0_6"])
    ws.cell(row=row, column=start_col + 8, value=vaccine_data["distance_6_12"])


def write_fda_header_row(ws, row: int, header_font: Font):
    write_reference_vaccine_header_row(ws, row, header_font, FDA_EXCEL_COLUMNS)


def write_fda_row(ws, row: int, fda_data: dict):
    write_reference_vaccine_row(ws, row, fda_data)


def write_pfizer_header_row(ws, row: int, header_font: Font):
    write_reference_vaccine_header_row(ws, row, header_font, PFIZER_EXCEL_COLUMNS)


def write_pfizer_row(ws, row: int, pfizer_data: dict):
    write_reference_vaccine_row(ws, row, pfizer_data)


def write_oracle_header_row(ws, row: int, header_font: Font):
    start_col = block_start_column(0)
    for offset, header in enumerate(ORACLE_EXCEL_COLUMNS):
        cell = ws.cell(row=row, column=start_col + offset, value=header)
        cell.font = header_font


def write_oracle_k_block(
    ws,
    first_row: int,
    last_row: int,
    oracle_rows: List[dict],
):
    if not oracle_rows:
        return

    start_col = block_start_column(0)
    shared = oracle_rows[0]
    merge_cols = [0, 6, 7, 8]

    for col_offset in merge_cols:
        col = start_col + col_offset
        if last_row > first_row:
            ws.merge_cells(
                start_row=first_row,
                start_column=col,
                end_row=last_row,
                end_column=col,
            )
        value = {
            0: shared["fit_slice_label"],
            6: shared["distance_current"],
            7: shared["distance_0_6"],
            8: shared["distance_6_12"],
        }[col_offset]
        cell = ws.cell(row=first_row, column=col, value=value)
        cell.alignment = Alignment(vertical="center")

    for row_offset, oracle_row in enumerate(oracle_rows):
        row = first_row + row_offset
        ws.cell(row=row, column=start_col + 2, value=oracle_row["medoid"])
        ws.cell(row=row, column=start_col + 3, value=oracle_row["lineage"])
        ws.cell(row=row, column=start_col + 4, value=oracle_row["covfit_fitness"])
        ws.cell(row=row, column=start_col + 5, value=oracle_row["pyro_fitness"])


def write_oracle_sections(
    ws,
    start_row: int,
    cutoff_oracle: Dict[str, Dict[int, List[dict]]],
    total_cols: int,
    k_fill: PatternFill,
    oracle_fill: PatternFill,
) -> int:
    current_row = start_row

    for oracle_interval in DISJOINT_WINDOWS:
        interval_rows = cutoff_oracle.get(oracle_interval, {})
        if not interval_rows:
            continue

        fit_label = ORACLE_INTERVAL_LABELS[oracle_interval]
        ws.cell(row=current_row, column=1, value=f"Oracle Medoids (fit: {fit_label})")
        ws.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=total_cols,
        )
        banner = ws.cell(row=current_row, column=1)
        banner.fill = oracle_fill
        banner.font = Font(bold=True)
        banner.alignment = Alignment(horizontal="left")
        current_row += 2

        for k in K_VALUES:
            k_rows = interval_rows.get(k, [])
            if not k_rows:
                continue

            ws.cell(row=current_row, column=1, value=f"K={k}")
            ws.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=total_cols,
            )
            k_banner = ws.cell(row=current_row, column=1)
            k_banner.fill = k_fill
            k_banner.font = Font(bold=True)
            k_banner.alignment = Alignment(horizontal="left")
            current_row += 1

            write_oracle_header_row(ws, current_row, Font(bold=True))
            current_row += 1

            first_data_row = current_row
            last_data_row = current_row + k - 1
            write_oracle_k_block(ws, first_data_row, last_data_row, k_rows)
            current_row = last_data_row + 2

    return current_row


def write_excel_report(
    all_rows: List[dict],
    fda_results: Dict[str, dict],
    pfizer_results: Dict[str, dict],
    oracle_results: Dict[str, Dict[str, Dict[int, List[dict]]]],
    output_path: Path,
):
    wb = Workbook()
    wb.remove(wb.active)

    date_fill = PatternFill("solid", fgColor="BDD7EE")
    k_fill = PatternFill("solid", fgColor="FFF2CC")
    fda_fill = PatternFill("solid", fgColor="E2EFDA")
    pfizer_fill = PatternFill("solid", fgColor="DDEBF7")
    oracle_fill = PatternFill("solid", fgColor="FCE4D6")
    header_font = Font(bold=True)

    nested = nest_medoid_rows(all_rows)
    cutoff_dates = sorted(
        set(nested) | set(fda_results) | set(pfizer_results) | set(oracle_results)
    )

    for cutoff_date in cutoff_dates:
        sheet_name = cutoff_display_name(cutoff_date)
        ws = wb.create_sheet(title=sheet_name[:31])
        sheet_rows = nested.get(cutoff_date, {})
        betas = sorted({beta for k_rows in sheet_rows.values() for beta in k_rows})
        if not betas:
            betas = sorted(BETAS_DEFAULT)
        total_cols = sheet_total_columns(betas)

        ws.cell(row=1, column=1, value=cutoff_display_name(cutoff_date))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        ws["A1"].fill = date_fill
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center")

        fda_data = fda_results.get(cutoff_date)
        fda_strain = FDA_STRAINS_BY_CUTOFF.get(cutoff_date, "")
        ws.cell(row=2, column=1, value=f"FDA Recommended Variant: {fda_strain}")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
        fda_banner = ws.cell(row=2, column=1)
        fda_banner.fill = fda_fill
        fda_banner.font = Font(bold=True)
        fda_banner.alignment = Alignment(horizontal="left")

        current_row = 4
        write_fda_header_row(ws, current_row, header_font)
        current_row += 1

        if fda_data:
            write_fda_row(ws, current_row, fda_data)
        current_row += 2

        pfizer_data = pfizer_results.get(cutoff_date)
        pfizer_inn = PFIZER_INN_BY_CUTOFF.get(cutoff_date, "")
        ws.cell(row=current_row, column=1, value=f"Pfizer Vaccine: {pfizer_inn}")
        ws.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=total_cols,
        )
        pfizer_banner = ws.cell(row=current_row, column=1)
        pfizer_banner.fill = pfizer_fill
        pfizer_banner.font = Font(bold=True)
        pfizer_banner.alignment = Alignment(horizontal="left")
        current_row += 2

        write_pfizer_header_row(ws, current_row, header_font)
        current_row += 1

        if pfizer_data:
            write_pfizer_row(ws, current_row, pfizer_data)
        current_row += 2

        cutoff_oracle = oracle_results.get(cutoff_date, {})
        if cutoff_oracle:
            current_row = write_oracle_sections(
                ws,
                current_row,
                cutoff_oracle,
                total_cols,
                k_fill,
                oracle_fill,
            )

        for k in K_VALUES:
            if k not in sheet_rows:
                continue

            ws.cell(row=current_row, column=1, value=f"K={k}")
            ws.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=total_cols,
            )
            banner = ws.cell(row=current_row, column=1)
            banner.fill = k_fill
            banner.font = Font(bold=True)
            banner.alignment = Alignment(horizontal="left")
            current_row += 1

            write_beta_header_row(ws, current_row, betas, header_font)
            current_row += 1

            first_data_row = current_row
            last_data_row = current_row + k - 1

            for beta_idx, beta in enumerate(betas):
                k_rows = sheet_rows.get(k, {}).get(beta, [])
                write_medoid_block(
                    ws,
                    first_data_row,
                    last_data_row,
                    block_start_column(beta_idx),
                    k_rows,
                    beta,
                )

            current_row = last_data_row + 2

        block_widths = [16, 28, 14, 18, 16, 14, 18, 28, 28]
        for col_idx in range(1, total_cols + 1):
            block_offset = (col_idx - 1) % (BLOCK_COLS + BLOCK_GUTTER)
            if block_offset < BLOCK_COLS:
                ws.column_dimensions[get_column_letter(col_idx)].width = block_widths[block_offset]
            else:
                ws.column_dimensions[get_column_letter(col_idx)].width = 2

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"\nWrote Excel report: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Export scout medoid analysis to Excel.")
    parser.add_argument("--root-path", default=ROOT_PATH_DEFAULT)
    parser.add_argument("--output", default=str(REPO_ROOT / "scout_medoid_report.xlsx"))
    parser.add_argument("--cutoff-dates", nargs="+", default=CUTOFF_DATES_DEFAULT)
    parser.add_argument("--betas", nargs="+", type=float, default=BETAS_DEFAULT)
    parser.add_argument(
        "--skip-covfit",
        action="store_true",
        help="Reuse existing CoVFit TSVs when present; missing predictions are generated automatically.",
    )
    parser.add_argument("--skip-pairwise", action="store_true", help="Reuse existing pairwise_dist.parquet.")
    parser.add_argument(
        "--extract-lineage-only",
        action="store_true",
        help="Parse existing lineage logs instead of regenerating them.",
    )
    args = parser.parse_args()

    root_path = Path(args.root_path)
    output_path = Path(args.output)
    covfit_dir = find_covfit_dir()

    print("Loading shared data...")
    final_df = pd.read_parquet(root_path / "final_df.parquet")
    id_to_seq = load_data(root_path / "sequence_dict.parquet")
    ref_seq = getReference(str(root_path / "ref_spike_genome.fa"))
    pyro_df = read_pickle(root_path / f"pyro_ver/{PYRO_VER}/PyroSpike.pkl")
    mutation_score_dict = pyro_df.set_index("mutation")["R / R_A"].to_dict()
    lineage_fasta = root_path / "lineage_spike.fasta"
    lineage_groups = load_lineage_spike_groups(lineage_fasta)

    all_rows: List[dict] = []
    fda_results: Dict[str, dict] = {}
    pfizer_results: Dict[str, dict] = {}
    oracle_results: Dict[str, Dict[str, Dict[int, List[dict]]]] = {}

    for cutoff_date in args.cutoff_dates:
        (
            current_seq_df,
            future_seq_dfs,
            dist_matrix,
            weights,
            _id2idx,
            idx2id,
            fitness_scores,
            dist_df,
        ) = prepare_cutoff_assets(
            root_path=root_path,
            cutoff_date=cutoff_date,
            final_df=final_df,
            id_to_seq=id_to_seq,
            covfit_dir=covfit_dir,
            skip_pairwise=args.skip_pairwise,
        )

        clade_to_seq, vaccine_fitness = prepare_vaccine_covfit(
            root_path=root_path,
            cutoff_date=cutoff_date,
            covfit_dir=covfit_dir,
        )
        fda_strain = FDA_STRAINS_BY_CUTOFF.get(cutoff_date)
        if fda_strain:
            fda_results[cutoff_date] = run_fda_variant_analysis(
                strain=fda_strain,
                clade_to_seq=clade_to_seq,
                vaccine_fitness=vaccine_fitness,
                current_seq_df=current_seq_df,
                future_seq_dfs=future_seq_dfs,
                id_to_seq=id_to_seq,
                ref_seq=ref_seq,
                mutation_score_dict=mutation_score_dict,
                lineage_groups=lineage_groups,
            )
            print(
                f"  FDA variant {fda_strain}: "
                f"current={fda_results[cutoff_date]['distance_current']:.3f}, "
                f"0-6mo={fda_results[cutoff_date]['distance_0_6']:.3f}, "
                f"6-12mo={fda_results[cutoff_date]['distance_6_12']:.3f}"
            )

        if cutoff_date in PFIZER_INN_BY_CUTOFF:
            pfizer_inn, pfizer_seq = load_pfizer_sequence(cutoff_date)
            pfizer_fitness = prepare_pfizer_covfit(
                root_path=root_path,
                cutoff_date=cutoff_date,
                covfit_dir=covfit_dir,
                pfizer_inn=pfizer_inn,
                pfizer_seq=pfizer_seq,
            )
            pfizer_results[cutoff_date] = run_pfizer_variant_analysis(
                pfizer_inn=pfizer_inn,
                pfizer_seq=pfizer_seq,
                pfizer_fitness=pfizer_fitness,
                current_seq_df=current_seq_df,
                future_seq_dfs=future_seq_dfs,
                id_to_seq=id_to_seq,
                ref_seq=ref_seq,
                mutation_score_dict=mutation_score_dict,
                lineage_groups=lineage_groups,
            )
            print(
                f"  Pfizer vaccine {pfizer_inn}: "
                f"current={pfizer_results[cutoff_date]['distance_current']:.3f}, "
                f"0-6mo={pfizer_results[cutoff_date]['distance_0_6']:.3f}, "
                f"6-12mo={pfizer_results[cutoff_date]['distance_6_12']:.3f}"
            )

        oracle_results[cutoff_date] = run_oracle_analysis(
            cutoff_date=cutoff_date,
            current_seq_df=current_seq_df,
            future_seq_dfs=future_seq_dfs,
            dist_df=dist_df,
            id_to_seq=id_to_seq,
            ref_seq=ref_seq,
            mutation_score_dict=mutation_score_dict,
            lineage_groups=lineage_groups,
        )

        for beta in args.betas:
            rows = run_beta_analysis(
                beta=beta,
                dist_matrix=dist_matrix,
                weights=weights,
                idx2id=idx2id,
                fitness_scores=fitness_scores,
                current_seq_df=current_seq_df,
                future_seq_dfs=future_seq_dfs,
                id_to_seq=id_to_seq,
                ref_seq=ref_seq,
                mutation_score_dict=mutation_score_dict,
                lineage_groups=lineage_groups,
                lineage_fasta=lineage_fasta,
                root_path=root_path,
                cutoff_date=cutoff_date,
                write_lineage=not args.extract_lineage_only,
            )
            all_rows.extend(rows)

    write_excel_report(all_rows, fda_results, pfizer_results, oracle_results, output_path)

    summary_df = pd.DataFrame(all_rows)
    oracle_flat_rows = [
        row
        for cutoff_data in oracle_results.values()
        for interval_data in cutoff_data.values()
        for k_rows in interval_data.values()
        for row in k_rows
    ]
    if oracle_flat_rows:
        oracle_df = pd.DataFrame(oracle_flat_rows)
        summary_df = pd.concat([summary_df, oracle_df], ignore_index=True, sort=False)
    if fda_results:
        fda_df = pd.DataFrame(
            [
                {"cutoff_date": cutoff, "reference_type": "FDA", **data}
                for cutoff, data in fda_results.items()
            ]
        )
        summary_df = pd.concat([fda_df, summary_df], ignore_index=True, sort=False)
    if pfizer_results:
        pfizer_df = pd.DataFrame(
            [
                {"cutoff_date": cutoff, "reference_type": "Pfizer", **data}
                for cutoff, data in pfizer_results.items()
            ]
        )
        summary_df = pd.concat([pfizer_df, summary_df], ignore_index=True, sort=False)
    summary_csv = output_path.with_suffix(".csv")
    summary_df.to_csv(summary_csv, index=False)
    print(f"Wrote CSV summary: {summary_csv}")


if __name__ == "__main__":
    main()
